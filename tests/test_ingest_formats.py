"""Tests du support multi-formats de storage/ingest.py (txt, md, pdf, docx).

Couvre : iter_text, read_blocks (nouvelle version par accumulation),
SUPPORTED_EXTENSIONS, le filtrage par extension dans diff(), et l'ingestion
bout en bout de vrais fichiers PDF et DOCX.

Dépendances : pip install pytest pytest-asyncio pypdf python-docx
(les tests PDF / DOCX sont ignorés automatiquement si la lib manque).
Le PDF de test est généré à la main : pas besoin de reportlab.
"""

from pathlib import Path

import pytest
from sqlmodel import Session, SQLModel, create_engine, select

from config.db.models import IngestedFile
from storage import ingest as ingest_mod
from storage.ingest import SUPPORTED_EXTENSIONS, diff, file_hash, iter_text, read_blocks


# --------------------------------------------------------------------------- #
# Générateurs de fichiers de test
# --------------------------------------------------------------------------- #
def make_docx(path: Path, paragraphs: list[str], table: list[list[str]] | None = None) -> Path:
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    doc.save(str(path))
    return path


def make_pdf(path: Path, pages: list[str]) -> Path:
    """Écrit un PDF minimal valide (police Helvetica, une ligne de texte par page).
    Le texte doit être en ASCII."""
    pytest.importorskip("pypdf")

    n = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {n} >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {5 + 2 * i} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        )
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET"
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")

    out = b"%PDF-1.4\n"
    offsets = []
    for num, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{num} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    path.write_bytes(out)
    return path


def write(path: Path, content: str) -> Path:
    path.write_text(content, encoding="utf-8")
    return path


def full_text(path: Path) -> str:
    return "".join(iter_text(str(path)))


# --------------------------------------------------------------------------- #
# Fakes / fixtures (mêmes principes que test_ingest.py)
# --------------------------------------------------------------------------- #
class FakeEmbedder:
    async def embed(self, texts: list[str]) -> list[list[float]]:
        return [[float(len(t))] for t in texts]


class FakeChunkStore:
    def __init__(self) -> None:
        self.deleted: list[int] = []
        self.added: list[tuple[list, list]] = []

    async def delete_by_source(self, source_id: int) -> None:
        self.deleted.append(source_id)

    async def add_chunks(self, chunks: list, embeddings: list) -> None:
        self.added.append((chunks, embeddings))

    def all_chunks(self) -> list:
        return [c for chunks, _ in self.added for c in chunks]


@pytest.fixture
def engine(tmp_path, monkeypatch):
    eng = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    SQLModel.metadata.create_all(eng)
    monkeypatch.setattr(ingest_mod, "engine", eng)
    return eng


@pytest.fixture
def docs(tmp_path) -> Path:
    d = tmp_path / "docs"
    d.mkdir()
    return d


@pytest.fixture
def store(monkeypatch) -> FakeChunkStore:
    fake = FakeChunkStore()
    monkeypatch.setattr(ingest_mod, "SqliteVectorStorage", lambda engine: object())
    monkeypatch.setattr(ingest_mod, "ChunkStore", lambda engine, vector_store: fake)
    monkeypatch.setattr(ingest_mod, "OllamaEmbedder", FakeEmbedder)
    return fake


def all_rows(engine) -> list[IngestedFile]:
    with Session(engine) as s:
        return list(s.exec(select(IngestedFile)).all())


# --------------------------------------------------------------------------- #
# SUPPORTED_EXTENSIONS
# --------------------------------------------------------------------------- #
def test_supported_extensions_contents():
    assert {".txt", ".md", ".pdf", ".docx"} <= SUPPORTED_EXTENSIONS
    assert ".png" not in SUPPORTED_EXTENSIONS


# --------------------------------------------------------------------------- #
# iter_text : texte brut
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("name", ["a.txt", "a.md"])
def test_iter_text_plain_text_formats(tmp_path, name):
    f = write(tmp_path / name, "héllo wörld\nligne 2")
    assert full_text(f) == "héllo wörld\nligne 2"


def test_iter_text_extension_is_case_insensitive(tmp_path):
    f = write(tmp_path / "A.TXT", "majuscules")
    assert full_text(f) == "majuscules"


def test_iter_text_large_text_file_is_streamed_in_pieces(tmp_path):
    content = "x" * 70_000
    f = write(tmp_path / "big.txt", content)

    pieces = list(iter_text(str(f)))

    assert len(pieces) == 2  # 65536 + le reste
    assert "".join(pieces) == content


def test_iter_text_empty_text_file(tmp_path):
    f = write(tmp_path / "empty.txt", "")
    assert list(iter_text(str(f))) == []


def test_iter_text_unsupported_extension_raises(tmp_path):
    f = tmp_path / "image.png"
    f.write_bytes(b"\x89PNG")
    with pytest.raises(ValueError, match="Unsupported"):
        list(iter_text(str(f)))


# --------------------------------------------------------------------------- #
# iter_text : DOCX
# --------------------------------------------------------------------------- #
def test_iter_text_docx_paragraphs(tmp_path):
    f = make_docx(tmp_path / "a.docx", ["Premier paragraphe", "Deuxième paragraphe"])

    text = full_text(f)

    assert "Premier paragraphe\nDeuxième paragraphe\n" in text


def test_iter_text_docx_table_rows_are_pipe_separated(tmp_path):
    f = make_docx(
        tmp_path / "t.docx",
        ["Intro"],
        table=[["nom", "age"], ["alice", "30"]],
    )

    text = full_text(f)

    assert "nom | age\n" in text
    assert "alice | 30\n" in text


def test_iter_text_docx_tables_come_after_paragraphs(tmp_path):
    f = make_docx(tmp_path / "t.docx", ["Paragraphe"], table=[["cellule"]])

    text = full_text(f)

    assert text.index("Paragraphe") < text.index("cellule")


# --------------------------------------------------------------------------- #
# iter_text : PDF
# --------------------------------------------------------------------------- #
def test_iter_text_pdf_yields_one_piece_per_page(tmp_path):
    f = make_pdf(tmp_path / "a.pdf", ["Hello from page one", "Hello from page two"])

    pieces = list(iter_text(str(f)))

    assert len(pieces) == 2
    assert "page one" in pieces[0]
    assert "page two" in pieces[1]
    assert all(p.endswith("\n") for p in pieces)


def test_iter_text_pdf_page_without_text_gives_only_newline(tmp_path):
    f = make_pdf(tmp_path / "blank.pdf", ["", "Texte page deux"])

    pieces = list(iter_text(str(f)))

    assert pieces[0].strip() == ""
    assert "Texte page deux" in pieces[1]


# --------------------------------------------------------------------------- #
# read_blocks : logique d'accumulation
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("length", "expected_blocks"),
    [(0, 0), (1, 1), (10, 1), (11, 2), (17, 2), (18, 3)],
)
def test_read_blocks_tail_is_never_a_pure_overlap(tmp_path, length, expected_blocks):
    """block_size=10, overlap=3 : le reste n'est émis que s'il apporte du texte neuf."""
    f = write(tmp_path / "a.txt", "abcdefghijklmnopqrstuvwxyz"[:length])

    blocks = list(read_blocks(str(f), block_size=10, overlap=3))

    assert len(blocks) == expected_blocks


@pytest.mark.parametrize("piece_size", [1, 7, 100, 5000])
def test_read_blocks_result_does_not_depend_on_piece_size(monkeypatch, piece_size):
    """Que iter_text renvoie le texte en miettes ou d'un bloc, les blocs sont identiques."""
    text = "".join(chr(97 + i % 26) for i in range(1234))

    monkeypatch.setattr(
        ingest_mod,
        "iter_text",
        lambda path: (text[i : i + piece_size] for i in range(0, len(text), piece_size)),
    )
    blocks = list(read_blocks("ignored", block_size=100, overlap=15))

    monkeypatch.setattr(ingest_mod, "iter_text", lambda path: iter([text]))
    reference = list(read_blocks("ignored", block_size=100, overlap=15))

    assert blocks == reference


def test_read_blocks_unsupported_extension_raises_on_iteration(tmp_path):
    f = tmp_path / "image.png"
    f.write_bytes(b"\x89PNG")
    with pytest.raises(ValueError):
        list(read_blocks(str(f)))


# --------------------------------------------------------------------------- #
# read_blocks : sur de vrais DOCX / PDF
# --------------------------------------------------------------------------- #
def _check_blocks_are_consistent(path: Path, block_size: int, overlap: int) -> list[str]:
    blocks = [b for b, _ in read_blocks(str(path), block_size, overlap)]
    assert len(blocks) > 1
    assert all(len(b) <= block_size for b in blocks)
    for prev, nxt in zip(blocks, blocks[1:], strict=False):
        assert prev[-overlap:] == nxt[:overlap]
    rebuilt = blocks[0] + "".join(b[overlap:] for b in blocks[1:])
    assert rebuilt == full_text(path)
    return blocks


def test_read_blocks_on_docx(tmp_path):
    paragraphs = [f"Paragraphe numéro {i} avec un peu de texte" for i in range(30)]
    f = make_docx(tmp_path / "long.docx", paragraphs)

    _check_blocks_are_consistent(f, block_size=100, overlap=10)


def test_read_blocks_on_pdf(tmp_path):
    f = make_pdf(tmp_path / "long.pdf", ["lorem ipsum " * 30] * 3)

    _check_blocks_are_consistent(f, block_size=100, overlap=10)


def test_read_blocks_indexes_are_sequential_on_docx(tmp_path):
    f = make_docx(tmp_path / "long.docx", ["texte " * 50] * 5)

    indexes = [i for _, i in read_blocks(str(f), block_size=100, overlap=10)]

    assert indexes == list(range(len(indexes)))


# --------------------------------------------------------------------------- #
# diff : filtrage par extension
# (échoue tant que diff() ne filtre pas sur SUPPORTED_EXTENSIONS)
# --------------------------------------------------------------------------- #
def test_diff_ignores_unsupported_extensions(engine, docs):
    kept = [
        write(docs / "a.txt", "a"),
        write(docs / "b.md", "b"),
        docs / "c.pdf",
        docs / "d.docx",
        docs / "E.PDF",  # extension en majuscules
    ]
    kept[2].write_bytes(b"%PDF fake")
    kept[3].write_bytes(b"PK fake")
    kept[4].write_bytes(b"%PDF fake 2")
    (docs / "image.png").write_bytes(b"\x89PNG")
    (docs / "archive.zip").write_bytes(b"PK")
    (docs / "noext").write_bytes(b"data")

    to_remove, to_add = diff(str(docs))

    assert to_remove == []
    assert {p.path for p in to_add} == {p.as_posix() for p in kept}


def test_diff_supported_file_in_subfolder_is_found(engine, docs):
    (docs / "sub").mkdir()
    f = write(docs / "sub" / "deep.md", "profond")

    _, to_add = diff(str(docs))

    assert [p.path for p in to_add] == [f.as_posix()]
    assert to_add[0].hash == file_hash(f)


# --------------------------------------------------------------------------- #
# ingest bout en bout avec de vrais PDF / DOCX
# --------------------------------------------------------------------------- #
@pytest.mark.asyncio
async def test_ingest_indexes_pdf_and_docx_content(engine, docs, store):
    make_docx(docs / "note.docx", ["Le docx parle de bananes"])
    make_pdf(docs / "rapport.pdf", ["Le pdf parle de pommes"])
    write(docs / "readme.md", "Le markdown parle de poires")

    await ingest_mod.ingest(str(docs))

    assert {r.file_name for r in all_rows(engine)} == {"note.docx", "rapport.pdf", "readme.md"}
    all_text = " ".join(c.text for c in store.all_chunks())
    assert "bananes" in all_text
    assert "pommes" in all_text
    assert "poires" in all_text


@pytest.mark.asyncio
async def test_ingest_skips_unsupported_files(engine, docs, store):
    write(docs / "a.txt", "contenu utile " * 10)
    (docs / "image.png").write_bytes(b"\x89PNG not text")
    (docs / "archive.zip").write_bytes(b"PK\x03\x04")

    await ingest_mod.ingest(str(docs))

    assert [r.file_name for r in all_rows(engine)] == ["a.txt"]


@pytest.mark.asyncio
async def test_ingest_modified_docx_is_reindexed(engine, docs, store):
    f = make_docx(docs / "note.docx", ["version un du document"])
    await ingest_mod.ingest(str(docs))
    old = all_rows(engine)[0]
    store.added.clear()

    make_docx(docs / "note.docx", ["version deux du document, modifiée"])
    await ingest_mod.ingest(str(docs))

    rows = all_rows(engine)
    assert len(rows) == 1
    assert rows[0].content_hash == file_hash(f) != old.content_hash
    assert store.deleted == [old.id]
    assert "version deux" in " ".join(c.text for c in store.all_chunks())
